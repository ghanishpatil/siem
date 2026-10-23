# Generates the SentinelX ICS project report (.docx) in the department format.
import os

from docx import Document
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

TMP = os.environ.get("TEMP", ".")
FIG31 = os.path.join(TMP, "siem_fig31.png")
EVID = r"f:\siem\docs\phase7-browser-evidence"
OUT = r"f:\siem\docs\SentinelX_ICS_Project_Report.docx"

BLACK = RGBColor(0, 0, 0)


# ---------------------------------------------------------------- Figure 3.1
def _font(size, bold=False):
    for name in ("arialbd.ttf" if bold else "arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            continue
    return ImageFont.load_default()


def draw_architecture(path):
    W, H = 1480, 340
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    f = _font(20)
    fb = _font(21, bold=True)
    boxes = [
        ("Event Sources", "logs / Simulation Lab"),
        ("Ingestion &", "Normalization (API)"),
        ("SQLite", "Store"),
        ("Correlation Engine", "5 rules + SHA-256"),
        ("Alerts &", "Evidence"),
        ("Dashboard", "Analytics / Export"),
    ]
    n = len(boxes)
    bw, bh = 200, 100
    gap = (W - 40 - n * bw) / (n - 1)
    y = 90
    centers = []
    for i, (l1, l2) in enumerate(boxes):
        x = 20 + i * (bw + gap)
        d.rounded_rectangle([x, y, x + bw, y + bh], radius=10, outline="black", width=3, fill=(245, 245, 245))
        for j, line in enumerate((l1, l2)):
            fnt = fb if j == 0 else f
            tw = d.textlength(line, font=fnt)
            d.text((x + (bw - tw) / 2, y + 26 + j * 30), line, fill="black", font=fnt)
        centers.append((x + bw, x, y + bh / 2))
    # arrows between boxes
    for i in range(n - 1):
        x_end = centers[i][0]
        x_next = centers[i + 1][1]
        cy = int(centers[i][2])
        d.line([x_end + 4, cy, x_next - 8, cy], fill="black", width=3)
        d.polygon([(x_next - 8, cy - 6), (x_next - 8, cy + 6), (x_next, cy)], fill="black")
    # title strip
    t = "Deterministic, explicit-correlation pipeline (no machine learning)"
    d.text((20, 250), t, fill=(70, 70, 70), font=_font(18))
    d.text((20, 282), "Frontend: React + Vite  |  Backend: FastAPI + SQLAlchemy  |  Storage: SQLite",
           fill=(70, 70, 70), font=_font(18))
    img.save(path)


# ---------------------------------------------------------------- doc helpers
doc = Document()

sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.left_margin = Inches(1.25)
sec.right_margin = Inches(1.0)
sec.top_margin = Inches(1.0)
sec.bottom_margin = Inches(1.0)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.font.color.rgb = BLACK
_rpr = normal.element.get_or_add_rPr()
_rf = _rpr.get_or_add_rFonts()
for a in ("w:ascii", "w:hAnsi", "w:cs"):
    _rf.set(qn(a), "Times New Roman")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def para(text="", align="justify", bold=False, size=12, italic=False, font="Times New Roman",
         space_after=6, line_spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = ALIGN[align]
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = BLACK
    return p


def chapter(text):
    para(text, align="left", bold=True, size=14, space_after=6, line_spacing=1.5)


def heading(text):  # unnumbered section heading (Abstract, References)
    para(text, align="left", bold=True, size=14, space_after=6)


def code_block(lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln if ln else " ")
        r.font.name = "Courier New"
        r.font.size = Pt(9.5)
        r.font.color.rgb = BLACK


def caption(text):
    para(text, align="center", italic=True, size=11, space_after=10)


def pagebreak():
    doc.add_page_break()


# ---------------------------------------------------------------- COVER PAGE
para("A Project Report on", align="center", space_after=10)
para("SENTINELX: A DETERMINISTIC SIEM CORRELATION ENGINE", align="center", bold=True, size=16, space_after=14)
para("PS Code: CIS-117", align="center", bold=True, space_after=2)
para("Domain: [D__]: Security Operations and Threat Detection", align="center", space_after=2)
para("Complexity: Intermediate", align="center", space_after=14)
para("Submitted in partial fulfilment of the requirements of the course", align="center", space_after=2)
para("Cryptography and Information Security (24UCYPC301) Laboratory", align="center", space_after=2)
para("T.Y. B.Tech Cyber Security, Semester V", align="center", space_after=14)
para("Submitted by", align="center", space_after=2)
para("[Student Full Name]", align="center", bold=True, space_after=2)
para("PRN: [__________]     Roll No.: [__]", align="center", space_after=14)
para("Under the guidance of", align="center", space_after=2)
para("Dr Chetan Bawankar", align="center", bold=True, space_after=2)
para("Head of Department of Cyber Security", align="center", space_after=10)
para("Department of Cyber Security", align="center", space_after=2)
para("School of Engineering and Technology", align="center", space_after=2)
para("Sanjivani University, Kopargaon", align="center", space_after=2)
para("Academic Year 2026-2027", align="center", bold=True)
pagebreak()

# ---------------------------------------------------------------- CERTIFICATE
para("CERTIFICATE", align="center", bold=True, size=14, space_after=14)
para(
    'This is to certify that the project report entitled "SentinelX: A Deterministic SIEM '
    'Correlation Engine" (PS Code: CIS-117) submitted by [Student Full Name] (PRN: [__________]) '
    "is a bonafide record of the laboratory project work carried out by him/her under my guidance, "
    "in partial fulfilment of the requirements of the course Cryptography and Information Security "
    "(24UCYPC301) Laboratory, T.Y. B.Tech Cyber Security, during the Academic Year 2026-2027 at the "
    "Department of Cyber Security, School of Engineering and Technology, Sanjivani University, "
    "Kopargaon. The similarity index of this report is [__]%.",
    space_after=28,
)
para("(Signature)", align="left", space_after=2)
para("Mr Kunal Makar", align="left", space_after=2)
para("Subject Teacher", align="left", space_after=20)
para("(Signature)", align="right", space_after=2)
para("Dr. Chetan D. Bawankar", align="right", space_after=2)
para("Head, Department of Cyber Security", align="right", space_after=20)
para("Place: Kopargaon", align="left", space_after=2)
para("Date: __/__/____", align="left")
pagebreak()

# ---------------------------------------------------------------- REPORT BODY
para("SentinelX: A Deterministic SIEM Correlation Engine", align="center", bold=True, size=14, space_after=4)
para("Student: [Name]     PRN: [__________]     PS Code: CIS-117", align="center", space_after=12)

heading("Abstract")
para(
    "Security operations teams are overwhelmed by the volume of raw log events, in which the signal "
    "of a genuine attack is easily buried. SentinelX is a deterministic Security Information and "
    "Event Management (SIEM) correlation engine that ingests security events, normalizes them to a "
    "canonical schema, persists them, and correlates them against five explicit detection rules to "
    "produce explainable, evidence-linked alerts. Detection is rule-based and reproducible - there "
    "is no machine learning - and alert de-duplication uses a SHA-256 fingerprint of the rule and "
    "its correlation entities. The system is implemented with a FastAPI (Python) backend, a "
    "SQLAlchemy/SQLite data layer, and a React/TypeScript dashboard. It detects brute force, "
    "password spraying, failed-then-successful logins, port scans, and privilege escalation; it was "
    "validated by an automated suite of 70 passing tests and an end-to-end functional run. The "
    "result is a transparent, auditable detection pipeline suitable for a controlled demonstration.",
)
para("Keywords: SIEM, event correlation, intrusion detection, SHA-256 fingerprinting, deterministic "
     "rules, FastAPI", bold=False)

chapter("1. Introduction and Problem Statement")
para(
    "A SIEM centralises security-relevant events from many sources and helps analysts detect and "
    "investigate threats. As networks scale, the number of events grows far beyond what a human can "
    "read, and meaningful attacks appear only as patterns across many individual events rather than "
    "in any single line [1]. The problem addressed by this project is to correlate such events "
    "deterministically and to raise a small number of explainable alerts, each linked to the exact "
    "evidence that triggered it, so that an analyst can quickly understand and act on an incident.",
)
para(
    "Established tools exist - Splunk, the Elastic (ELK) stack, and open-source platforms such as "
    "Wazuh and OSSEC [5] - but they are heavyweight and their detection logic is often opaque to a "
    "learner. Attack techniques themselves are well catalogued by the MITRE ATT&CK framework [3], "
    "and secure log-management practice is standardised by NIST [1]. The gap this project addresses "
    "is a compact, fully transparent, deterministic correlation engine whose every alert can be "
    "justified exactly by its rule, threshold, time window, and evidence - built for study and "
    "demonstration rather than as a black box.",
)

chapter("2. Objectives")
para("1. Design and implement a deterministic SIEM pipeline that ingests, normalizes, and persists "
     "heterogeneous security events through a typed REST API.", bold=False)
para("2. Implement five explicit correlation rules covering credential attacks, reconnaissance, and "
     "privilege escalation, each producing explainable, evidence-linked alerts.", bold=False)
para("3. Guarantee alert integrity and prevent duplicate alerts using SHA-256 entity fingerprinting "
     "with idempotent upserts.", bold=False)
para("4. Provide an investigation and analytics dashboard (evidence timeline, status lifecycle, JSON "
     "export) and validate the system with an automated test suite.", bold=False)

chapter("3. Methodology and System Design")
para(
    "SentinelX follows a layered, data-flow architecture (Figure 3.1). Events enter through the "
    "ingestion API or the built-in Simulation Lab; each record is validated and normalized; valid "
    "records are stored in SQLite; and, only on an explicit request, the correlation engine "
    "evaluates the stored events against the active rules and upserts alerts with their evidence. "
    "The frontend is presentation-only and communicates with the backend solely over HTTP.",
)
if not os.path.exists(FIG31):
    draw_architecture(FIG31)
doc.add_picture(FIG31, width=Inches(6.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Figure 3.1: SentinelX system architecture and data-flow pipeline.")
para(
    "Modules. (a) Normalization validates each event - timestamps (ISO-8601 or epoch), IP addresses "
    "(via the standard ipaddress library), destination ports (0-65535) - and maps event types and "
    "severities to a controlled vocabulary; one bad record is rejected individually without failing "
    "the batch. (b) Persistence stores normalized events and deduplicated alerts. (c) The detection "
    "engine runs five pure rule functions over sliding time windows. (d) The service layer "
    "orchestrates loading, evaluation, and a single-transaction upsert of alerts.",
)
para(
    "Detection algorithm. Each rule groups events by a correlation entity (such as source IP or "
    "username) and counts occurrences inside a sliding window; a detection is raised when the count "
    "meets the rule threshold. The windowing is deterministic (Section 4). The cryptographic "
    "component of the system is a SHA-256 fingerprint computed over the rule identifier and the "
    "sorted correlation entities [2]; this fingerprint is the stable, collision-resistant key used "
    "to deduplicate alerts, so the same ongoing attack remains a single incident across re-runs.",
)
para("The five rules and their default parameters are:", bold=False)
para("- AUTH-001 Brute Force - 10 authentication failures from one source IP within 60 s (high).", bold=False, space_after=2)
para("- AUTH-002 Password Spraying - 5 distinct accounts from one source IP within 300 s (high).", bold=False, space_after=2)
para("- AUTH-003 Failed Logins then Success - 5 failures then a success for the same IP and account "
     "within 300 s (critical).", bold=False, space_after=2)
para("- NET-001 Port Scan - 15 distinct destination ports from one source IP within 60 s (medium).", bold=False, space_after=2)
para("- PRIV-001 Privilege Escalation - an escalation event followed by a sensitive action by the "
     "same user within 300 s (high).", bold=False, space_after=8)
para(
    "Threat model. The system is designed to detect the five attack patterns above and to keep each "
    "alert auditable through its stored evidence; SHA-256 fingerprinting protects the integrity of "
    "alert identity. Input-level threats (malformed, oversized, or duplicate events) are mitigated "
    "by strict validation, bounded ingestion limits, and event-UID de-duplication. A known "
    "limitation, analysed in Section 5, is the absence of an authentication layer on the API.",
)

chapter("4. Implementation")
para(
    "The backend uses Python 3.10.7 with FastAPI 0.143, SQLAlchemy 2.0, Uvicorn, and Pydantic 2 for "
    "validation; persistence is SQLite. The frontend uses React 19, Vite, TypeScript, and Tailwind "
    "CSS. Testing uses Pytest 9.1 with the FastAPI TestClient. The backend is organised into thin "
    "API routes, a normalization package, a pure detection package (engine plus five rules), and a "
    "repository layer that owns all database access; only the service layer touches the database "
    "during correlation, keeping the rules pure and independently testable.",
)
para("Snippet 4.1 - SHA-256 alert fingerprint (deterministic de-duplication key):", bold=True, space_after=2)
code_block([
    "import hashlib",
    "",
    "def make_fingerprint(rule_id, entities):",
    "    parts = [rule_id] + [f\"{k}={entities[k]}\" for k in sorted(entities)]",
    "    return hashlib.sha256(\"|\".join(parts).encode(\"utf-8\")).hexdigest()",
])
para("")
para("Snippet 4.2 - deterministic sliding-window counter (core of the threshold rules):", bold=True, space_after=2)
code_block([
    "def mark_count_windows(events, window_seconds, threshold):",
    "    window = timedelta(seconds=window_seconds)",
    "    marked, left = set(), 0",
    "    for right in range(len(events)):",
    "        while events[right].timestamp - events[left].timestamp > window:",
    "            left += 1",
    "        if right - left + 1 >= threshold:",
    "            marked.update(range(left, right + 1))",
    "    return [events[i] for i in sorted(marked)]",
])
para("")

chapter("5. Results and Security Analysis")
para(
    "The system was validated with an automated test suite and an end-to-end functional run against "
    "an isolated database. Table 5.1 lists representative test cases and their observed results.",
)
para("Table 5.1: Test Cases", align="center", bold=True, space_after=4)
rows = [
    ("Test ID", "Input / Action", "Expected Output", "Actual Output", "Status"),
    ("TC01", "Ingest 10 authentication_failure events (one IP, within 60 s)", "HTTP 201, 10 accepted, 0 rejected", "HTTP 201, accepted=10", "Pass"),
    ("TC02", "Run correlation on the brute-force events", "One AUTH-001 (high) alert created", "detections=1, created=1 (AUTH-001)", "Pass"),
    ("TC03", "Retrieve alert detail", "Alert links to its 10 evidence events", "10 evidence rows returned", "Pass"),
    ("TC04", "Re-run correlation on unchanged evidence", "No duplicate alert; existing unchanged", "created=0, unchanged alert", "Pass"),
    ("TC05", "PATCH an alert with an invalid status", "Rejected (HTTP 422)", "HTTP 422 enum error", "Pass"),
    ("TC06", "Ingest an event missing its timestamp", "Rejected with field error (HTTP 400)", "HTTP 400, field 'timestamp'", "Pass"),
    ("TC07", "Health check with database unavailable", "HTTP 503, database 'unavailable'", "HTTP 503, status 'degraded'", "Pass"),
]
table = doc.add_table(rows=len(rows), cols=5)
table.style = "Table Grid"
widths = [Inches(0.7), Inches(2.2), Inches(1.7), Inches(1.7), Inches(0.7)]
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = table.rows[ri].cells[ci]
        cell.width = widths[ci]
        cell.text = ""
        pcell = cell.paragraphs[0]
        pcell.paragraph_format.line_spacing = 1.0
        pcell.paragraph_format.space_after = Pt(0)
        r = pcell.add_run(val)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = (ri == 0)
        r.font.color.rgb = BLACK
para("")
para(
    "The complete backend suite reports 70 passing tests. A single correlation pass is bounded to at "
    "most 100,000 events to keep evaluation time predictable, and the Correlation Engine view "
    "reports per-run latency in milliseconds. The frontend production build (type-check plus Vite "
    "bundle) completes successfully. Figures 5.1-5.3 show representative dashboard output captured "
    "during an end-to-end run.",
)
shots = [
    ("01-analytics.png", "Figure 5.1: Security analytics - alerts by investigation status and events by provenance."),
    ("02-timeline.png", "Figure 5.2: Alert evidence timeline linking an incident to its triggering events."),
    ("03-status.png", "Figure 5.3: Investigation status update, confirmed and persisted by the backend."),
]
for fn, cap in shots:
    fp = os.path.join(EVID, fn)
    if os.path.exists(fp):
        doc.add_picture(fp, width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption(cap)
para(
    "Security analysis. Against the modelled attacks, the engine deterministically detects each "
    "pattern and attaches the exact evidence, so a detection can always be justified. Alert integrity "
    "is protected by SHA-256 fingerprinting, which yields a stable de-duplication key and preserves "
    "an analyst's status across re-runs (idempotent upsert). At the input layer, all fields are "
    "validated, batch and body sizes are bounded, duplicates are rejected by event UID, database "
    "access is through a parameterised ORM (mitigating SQL injection), and untrusted event fields "
    "are rendered as text in the UI (mitigating cross-site scripting). Error responses are generic, "
    "avoiding stack-trace disclosure, and the health endpoint reports a database outage honestly as "
    "HTTP 503. The principal limitation is that the API currently has no authentication or "
    "authorization; CORS restricts browser origins only and is not a substitute for authentication. "
    "The system must therefore be run as a controlled, access-restricted demonstration and not "
    "exposed publicly until an authentication layer is added.",
)

chapter("6. Conclusion and Future Scope")
para(
    "All four objectives were met: a deterministic ingestion-to-alert pipeline was built; five "
    "explicit correlation rules were implemented with explainable, evidence-linked alerts; alert "
    "integrity and de-duplication were achieved with SHA-256 fingerprinting and idempotent upserts; "
    "and an investigation and analytics dashboard was delivered and validated by 70 passing tests. "
    "The main limitations are the absence of authentication, the single-writer nature of SQLite, and "
    "the lack of automated backups. Future work includes adding authentication and authorization (or "
    "platform-level access control) before any public deployment, migrating to a client/server "
    "database such as PostgreSQL for horizontal scale, adding data retention and automated backups, "
    "and optionally complementing the deterministic rules with anomaly-based detection.",
)
para("GitHub Repository Link: [https://github.com/your-username/sentinelx-siem]", bold=False)

heading("References")
refs = [
    '[1] K. Kent and M. Souppaya, "Guide to Computer Security Log Management," NIST Special '
    'Publication 800-92, Sep. 2006.',
    '[2] National Institute of Standards and Technology, "Secure Hash Standard (SHS)," FIPS PUB '
    '180-4, Aug. 2015.',
    '[3] The MITRE Corporation, "MITRE ATT&CK Framework." [Online]. Available: '
    'https://attack.mitre.org',
    '[4] W. Stallings, Cryptography and Network Security: Principles and Practice, 8th ed. Hoboken, '
    'NJ, USA: Pearson, 2023.',
    '[5] S. Bhatt, P. K. Manadhata, and L. Zomlot, "The Operational Role of Security Information and '
    'Event Management Systems," IEEE Security & Privacy, vol. 12, no. 5, pp. 35-41, 2014.',
    '[6] J. Arkko, M. Cotton, and L. Vegoda, "IPv4 Address Blocks Reserved for Documentation," RFC '
    '5737, IETF, Jan. 2010.',
]
for r in refs:
    para(r, bold=False, space_after=4)

doc.save(OUT)
print("WROTE", OUT)
print("paragraphs:", len(doc.paragraphs))
